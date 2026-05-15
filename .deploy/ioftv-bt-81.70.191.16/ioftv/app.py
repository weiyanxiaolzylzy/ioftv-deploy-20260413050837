from fastapi import FastAPI, HTTPException
from fastapi.responses import Response
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from ifc_service import IFCService
from typing import List, Optional
import openpyxl
from copy import copy
from io import BytesIO
from urllib.parse import quote
import os
import re

app = FastAPI(title="IFC Backend Service", version="1.0.0")

# CORS 配置
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# 初始化 IFC 服务
ifc_service = IFCService(base_path="server/uploads/ifc")


class ExtractMarksRequest(BaseModel):
    filename: str


class ExtractElementsRequest(BaseModel):
    filename: str
    limit: int = 8000


@app.get("/")
async def root():
    return {"message": "IFC Backend Service", "version": "1.0.0"}


@app.post("/api/ifc/extract-marks")
async def extract_marks(request: ExtractMarksRequest):
    """
    从 IFC 文件提取构件编号映射
    """
    try:
        marks = ifc_service.extract_component_marks(request.filename)
        return {
            "success": True,
            "marks": marks,
            "count": len(marks)
        }
    except FileNotFoundError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"解析失败: {str(e)}")


class BuildIndexRequest(BaseModel):
    filename: str
    limit: int = 8000


@app.post("/api/ifc/elements")
async def extract_elements(request: ExtractElementsRequest):
    """
    批量提取构件列表
    """
    try:
        result = ifc_service.extract_elements(request.filename, request.limit)
        return result
    except FileNotFoundError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"解析失败: {str(e)}")


@app.post("/api/ifc/build-index")
async def build_index(request: BuildIndexRequest):
    """
    合并接口：一次性返回构件编号映射和构件列表（供前端 buildElementIndex 使用）
    """
    try:
        result = ifc_service.build_index(request.filename, request.limit)
        return result
    except FileNotFoundError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"解析失败: {str(e)}")


@app.get("/health")
async def health_check():
    """健康检查"""
    return {"status": "healthy"}


# ─── 质检表导出（openpyxl，保留原模板全部格式）────────────────────────

QC_TEMPLATE_DIR = os.path.join(os.path.dirname(__file__), "单个构件质检表")

# 每个模板的固定列映射（1-indexed）
# seq_col: 序号列, design_col: 设计尺寸列, self_col: 自检列, leader_col: 班组长列, remark_col: 备注列
# data_start: 数据起始行, group_cell: 加工班组单元格, sig_cells: 签名单元格
TEMPLATE_MAP = {
    "H型钢柱质检表Rev.2_2021.04.08.xlsx": {
        "seq_col": 1, "design_col": 6, "self_col": 7, "leader_col": 8, "remark_col": 9,
        "data_start": 7,
        "group_cell": (3, 7),
        "sig_cells": {"self_leader": (24, 3), "qa": (25, 3)},
    },
    "十字柱质检表Rev.1_2020.06.11.xlsx": {
        "seq_col": 1, "design_col": 6, "self_col": 7, "leader_col": 9, "remark_col": 11,
        "data_start": 7,
        "group_cell": (3, 7),
        "sig_cells": {"self_leader": (21, 3), "qa": (22, 3)},
    },
    "梁质检表Rev. 2_2021.04.08.xlsx": {
        "seq_col": 1, "design_col": 4, "self_col": 5, "leader_col": 6, "remark_col": 7,
        "data_start": 7,
        "group_cell": (3, 5),
        "sig_cells": {"self_leader": (31, 3), "qa": (33, 3)},
    },
    "箱型柱质检表Rev.3_2021.04.08.xlsx": {
        "seq_col": 1, "design_col": 6, "self_col": 7, "leader_col": 8, "remark_col": 9,
        "data_start": 7,
        "group_cell": (3, 7),
        "sig_cells": {"self_leader": (23, 3), "qa": (24, 3)},
    },
    "箱型梁质检表Rev.0_2024.09.13.xlsx": {
        "seq_col": 1, "design_col": 6, "self_col": 7, "leader_col": 8, "remark_col": 9,
        "data_start": 7,
        "group_cell": (3, 7),
        "sig_cells": {"self_leader": (19, 3), "qa": (20, 3)},
    },
    "钢管桁架质检表Rev.0_2020.06.18.xlsx": {
        "seq_col": 1, "design_col": 5, "self_col": 6, "leader_col": 9, "remark_col": 12,
        "data_start": 7,
        "group_cell": (3, 8),
        "sig_cells": {"self_leader": (24, 4), "qa": (26, 4)},
    },
    "无理论值圆管柱尺寸及焊缝质检表2023.10.25.xlsx": {
        "seq_col": None, "design_col": None, "self_col": None, "leader_col": None, "remark_col": None,
        "data_start": None,
        "group_cell": None,
        "sig_cells": {"self_leader": (27, 1)},
    },
}


class QcExportRow(BaseModel):
    seq: Optional[int] = None
    name: Optional[str] = None
    designValue: Optional[float] = None
    measuredValue: Optional[float] = None
    groupMeasuredValue: Optional[float] = None
    deviation: Optional[float] = None
    verdict: Optional[str] = None


class QcExportMeta(BaseModel):
    groupName: Optional[str] = ""
    selfInspectorName: Optional[str] = ""
    teamLeaderName: Optional[str] = ""
    qualityInspectorName: Optional[str] = ""


class QcExportRequest(BaseModel):
    componentNo: Optional[str] = ""
    rows: List[QcExportRow] = []
    meta: QcExportMeta = QcExportMeta()


@app.get("/api/qc-templates")
async def list_qc_templates():
    templates = []
    if not os.path.isdir(QC_TEMPLATE_DIR):
        return {"success": True, "data": []}
    for fname in sorted(os.listdir(QC_TEMPLATE_DIR)):
        if fname.endswith(".xlsx"):
            templates.append({"id": fname, "fileName": fname, "title": fname.rsplit(".", 1)[0]})
    return {"success": True, "data": templates}


@app.post("/api/qc-export/{template_id}")
async def export_qc(template_id: str, body: QcExportRequest):
    safe_name = os.path.basename(template_id)
    fp = os.path.join(QC_TEMPLATE_DIR, safe_name)
    if not os.path.isfile(fp):
        raise HTTPException(status_code=404, detail="模板不存在")

    tpl_cfg = TEMPLATE_MAP.get(safe_name)
    if not tpl_cfg:
        raise HTTPException(status_code=400, detail="未配置该模板的映射")

    wb = openpyxl.load_workbook(fp)
    ws = wb.active

    meta = body.meta
    component_no = body.componentNo or ""

    # ── 填写检验数据 ──
    if tpl_cfg["seq_col"] and tpl_cfg["data_start"]:
        by_seq = {}
        for r in body.rows:
            if r.seq and r.seq > 0:
                by_seq[r.seq] = r

        seq_col = tpl_cfg["seq_col"]
        for row_num in range(tpl_cfg["data_start"], ws.max_row + 1):
            cell_val = ws.cell(row_num, seq_col).value
            if cell_val is None:
                continue
            try:
                seq = int(str(cell_val).strip())
            except (ValueError, TypeError):
                continue

            payload = by_seq.get(seq)
            if not payload:
                continue

            if tpl_cfg["design_col"] and payload.designValue is not None:
                ws.cell(row_num, tpl_cfg["design_col"]).value = payload.designValue

            if tpl_cfg["self_col"] and payload.measuredValue is not None:
                ws.cell(row_num, tpl_cfg["self_col"]).value = payload.measuredValue

            if tpl_cfg["leader_col"] and payload.groupMeasuredValue is not None:
                ws.cell(row_num, tpl_cfg["leader_col"]).value = payload.groupMeasuredValue

    # ── 填写加工班组 ──
    if tpl_cfg["group_cell"] and meta.groupName:
        r, c = tpl_cfg["group_cell"]
        old = ws.cell(r, c).value or ""
        old_clean = str(old).replace(" ", "")
        if old_clean in ("加工班组", "施工班组", "班组", ""):
            label = "加工班组" if "加工" in old_clean else ("施工班组" if "施工" in old_clean else "加工班组")
            ws.cell(r, c).value = f"{label}：{meta.groupName}"
        else:
            ws.cell(r, c + 1).value = meta.groupName

    # ── 填写签名 ──
    sig = tpl_cfg.get("sig_cells", {})

    if "self_leader" in sig:
        r, c = sig["self_leader"]
        old = str(ws.cell(r, c).value or "")
        new_text = old
        if meta.selfInspectorName:
            new_text = re.sub(r"自检员[：:]\s*", f"自检员：{meta.selfInspectorName}  ", new_text)
        if meta.teamLeaderName:
            new_text = re.sub(r"班组长[：:]\s*", f"班组长：{meta.teamLeaderName}  ", new_text)
        ws.cell(r, c).value = new_text

    if "qa" in sig:
        r, c = sig["qa"]
        old = str(ws.cell(r, c).value or "")
        if meta.qualityInspectorName:
            new_text = re.sub(r"(质量检查员|质量检测员|质检员)[：:]\s*",
                             lambda m: f"{m.group(1)}：{meta.qualityInspectorName}  ", old)
            ws.cell(r, c).value = new_text

    # ── 输出 ──
    buf = BytesIO()
    wb.save(buf)
    wb.close()
    buf.seek(0)

    out_name = f"{component_no}_{safe_name}" if component_no else safe_name
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(out_name)}"},
    )


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8765)
